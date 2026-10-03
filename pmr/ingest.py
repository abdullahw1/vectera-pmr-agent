"""Content-based discovery and label-based parsers; never fixed workbook cells."""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

import fitz
import openpyxl
from .numeric import currency, finite_number, display


QUARTER_WORDS = {"first": 1, "second": 2, "third": 3, "fourth": 4}
ROMANS = dict(
    zip(
        "i ii iii iv v vi vii viii ix x xi xii xiii xiv xv xvi xvii xviii xix xx".split(),
        map(str, range(1, 21)),
    )
)
NUMBER_WORDS = dict(
    zip(
        "one two three four five six seven eight nine ten eleven twelve thirteen fourteen fifteen sixteen seventeen eighteen nineteen twenty".split(),
        map(str, range(1, 21)),
    )
)
ENTITY_ABBREVIATIONS = {"intl": "international", "mgmt": "management", "inv": "investment"}


def norm(text):
    text = str(text or "").lower().replace("%", " percent ").replace("$", " dollars ")
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def entity(text):
    text = re.sub(r"\bU\.\s*S\.", "US", str(text), flags=re.I)
    text = text.replace("&", " and ")
    tokens = norm(text).split()
    if tokens[-2:] == ["l", "p"]:
        tokens = tokens[:-2]
    while tokens and tokens[-1] in {"llc", "lp", "limited"}:
        tokens.pop()
    # Number words are vintage labels, not rewrites of asset names like One Financial Plaza.
    for i, token in enumerate(tokens):
        if token in NUMBER_WORDS and (
            (i and tokens[i - 1] in {"fund", "series", "vintage"})
            or (i == len(tokens) - 1 and "fund" in tokens[:i])
        ):
            tokens[i] = NUMBER_WORDS[token]
    return " ".join(ENTITY_ABBREVIATIONS.get(ROMANS.get(t, t), ROMANS.get(t, t)) for t in tokens)


def quarter(text):
    match = re.search(r"\b(First|Second|Third|Fourth) Quarter\s+(\d{4})", str(text), re.I)
    if match:
        return int(match[2]), QUARTER_WORDS[match[1].lower()]
    match = re.search(r"\b([1-4])\s*Q\s*(\d{2}|\d{4})\b", str(text), re.I)
    if match:
        year = int(match[2])
        return (year + 2000 if year < 100 else year, int(match[1]))
    return None


def bounds(q):
    year, number = q
    start = date(year, 3 * number - 2, 1)
    end = date(year + 1, 1, 1) if number == 4 else date(year, 3 * number + 1, 1)
    return start, end


def pdf_pages(path):
    with fitz.open(path) as doc:
        return [page.get_text(sort=True) for page in doc]


class Sheet:
    def __init__(self, path, worksheet, ledger):
        self.path, self.ws, self.ledger = path, worksheet, ledger

    def rows(self):
        return list(self.ws.iter_rows())

    def find(self, label, *, numeric=False):
        matches = [
            r
            for r in self.rows()
            if entity(r[0].value) == entity(label)
            and (not numeric or any(isinstance(c.value, (int, float)) for c in r[1:]))
        ]
        if len(matches) != 1:
            raise ValueError(f"{self.path.name}/{self.ws.title}: {label!r} matched {len(matches)} rows")
        return matches[0]

    def fact(self, cell):
        return self.ledger.add(
            cell.value, {"file": self.path.name, "sheet": self.ws.title, "cell": cell.coordinate}
        )

    def number(self, cell, *, blank_zero=False, monetary=False, metric=None, unit=None, entity_name=None):
        if cell.value is None and blank_zero:
            return currency(0) if monetary else 0.0, self.ledger.add(
                0,
                {
                    "file": self.path.name,
                    "sheet": self.ws.title,
                    "cell": cell.coordinate,
                    "interpretation": "blank capital flow = zero",
                },
            )
        if isinstance(cell.value, bool) or not isinstance(cell.value, (int, float)):
            raise ValueError(f"Expected number at {self.path.name}/{self.ws.title}/{cell.coordinate}")
        value = currency(cell.value) if monetary else finite_number(cell.value)
        source = {
            "file": self.path.name,
            "sheet": self.ws.title,
            "cell": cell.coordinate,
            "raw_value": str(cell.value),
        }
        unit = unit or ("USD" if monetary else "x" if "multiple" in norm(metric) else "%")
        metadata = dict(
            entity=entity_name or str(self.ws.cell(cell.row, 1).value),
            metric=metric,
            unit=unit,
            authority="source workbook",
            period=self.ledger.reporting_period,
            precision="cent" if monetary else "source precision",
            permitted_formats=[display(value, places) for places in [0, 1, 2]],
        )
        if monetary:
            metadata["permitted_formats"] += [
                "$" + display(value, 2),
                "$" + display(value / 1_000_000, 1) + "M",
            ]
        return value, self.ledger.add(value, source, numeric=metadata)

    def headers(self, marker="Investment", grouped=False):
        row = next((r for r in self.rows() if norm(r[0].value) == norm(marker)), None)
        if row is None:
            raise ValueError(f"Header {marker!r} missing in {self.ws.title}")
        result, group = {}, ""
        lower = list(self.ws.iter_rows(min_row=row[0].row + 1, max_row=row[0].row + 1))[0]
        for i, c in enumerate(row):
            if c.value is not None:
                group = str(c.value)
            key = norm(group + " " + str(lower[i].value or "")) if grouped and i else norm(c.value)
            if key:
                result[key] = i
        return row[0].row, result


def discover(root, client, requested, ledger):
    candidates = []
    books = {}
    for path in sorted(root.rglob("*.xlsx")):
        try:
            book = openpyxl.load_workbook(path, data_only=True)
        except Exception as error:
            ledger.discovery.append(
                dict(file=str(path.relative_to(root)), status="failed", reason=type(error).__name__)
            )
            ledger.issue(
                "unreadable_input",
                f"Cannot classify workbook {path.name}; exclude it explicitly or repair it",
                "blocker",
            )
            continue
        books[path] = book
        if "FundingStatus(Agg)" not in book.sheetnames:
            continue
        text = "\n".join(
            str(c.value) for row in book["FundingStatus(Agg)"] for c in row if isinstance(c.value, str)
        )
        labels = [
            str(c.value)
            for ws in book
            for row in ws
            for c in row
            if isinstance(c.value, str) and "benchmark" in c.value.lower()
        ]
        # The short code is present in the benchmark label; content outranks filenames.
        belongs = any(re.search(r"\b" + re.escape(client) + r"\b", label, re.I) for label in labels)
        if belongs and quarter(text) == requested:
            candidates.append((path, book))
        else:
            ledger.discovery.append(
                dict(
                    file=str(path.relative_to(root)),
                    status="ignored",
                    reason="Flash content has no requested client/quarter data",
                )
            )
    if len(candidates) != 1:
        raise ValueError(f"Expected one populated flash for {client}/{requested}, found {len(candidates)}")
    flash, book = candidates[0]
    ledger.discovery.append(
        dict(
            file=str(flash.relative_to(root)),
            status="accepted",
            reason="Populated flash client and period content match",
        )
    )
    ws = book["FundingStatus(Agg)"]
    # The workbook contract places its identity above the first labelled data section.
    identity = next(
        c for row in ws for c in row if isinstance(c.value, str) and not quarter(c.value) and c.value.strip()
    )
    name = re.sub(r"\s+Portfolio$", "", identity.value)
    name_id = Sheet(flash, ws, ledger).fact(identity)
    prior, managers, appendix, logs, history, deck = [], [], [], [], [], []
    for path in sorted(root.rglob("*")):
        if path.suffix.lower() == ".pdf":
            try:
                pages = pdf_pages(path)
            except Exception as error:
                ledger.discovery.append(
                    dict(file=str(path.relative_to(root)), status="failed", reason=type(error).__name__)
                )
                ledger.issue("unreadable_input", f"Cannot classify PDF {path.name}", "blocker")
                continue
            text = "\n".join(pages)
            period = quarter(text)
            belongs = norm(name) in norm(text)
            role = "ignored: different period/client or archived stub"
            if "Performance Measurement Report" in text and belongs:
                if period and period < requested:
                    prior.append((period, path, pages))
                    role = "prior PMR candidate"
            elif period == requested:
                if "Portfolio Composition" in text and belongs:
                    appendix.append(path)
                    role = "flash appendix candidate"
                else:
                    managers.append((path, pages))
                    role = "supporting document candidate; entity match pending"
            ledger.discovery.append(
                dict(
                    file=str(path.relative_to(root)),
                    status="ignored" if role.startswith("ignored") else "candidate",
                    reason=role,
                )
            )
        elif path.suffix.lower() == ".xlsx" and path != flash:
            other = books.get(path)
            if other is None:
                continue
            roles = []
            for sheet in other:
                strings = {norm(c.value) for row in sheet for c in row if isinstance(c.value, str)}
                if {"meeting date", "committee action", "client"} <= strings:
                    logs.append((path, sheet))
                    roles.append("committee log")
                if any("target allocation" in v for v in strings) and any(
                    "net asset value" in v for v in strings
                ):
                    history.append((path, sheet))
                    roles.append("allocation history")
            if roles:
                ledger.discovery.append(
                    dict(file=str(path.relative_to(root)), status="accepted", reason=", ".join(roles))
                )
            elif "FundingStatus(Agg)" not in other.sheetnames:
                ledger.discovery.append(
                    dict(
                        file=str(path.relative_to(root)),
                        status="ignored",
                        reason="No recognized source contract",
                    )
                )
        elif path.suffix.lower() == ".pptx":
            import zipfile
            import xml.etree.ElementTree as ET

            try:
                with zipfile.ZipFile(path) as archive:
                    text = " ".join(
                        " ".join(ET.fromstring(archive.read(n)).itertext())
                        for n in archive.namelist()
                        if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)
                    )
            except (zipfile.BadZipFile, ET.ParseError, OSError) as error:
                ledger.discovery.append(
                    dict(file=str(path.relative_to(root)), status="failed", reason=type(error).__name__)
                )
                ledger.issue("unreadable_input", f"Cannot classify deck {path.name}", "blocker")
                continue
            if quarter(text) == requested:
                deck.append(path)
            ledger.discovery.append(
                dict(
                    file=str(path.relative_to(root)),
                    status="candidate" if quarter(text) == requested else "ignored",
                    reason="Market deck period content",
                )
            )
        elif path.is_file() and path != flash:
            ledger.discovery.append(
                dict(file=str(path.relative_to(root)), status="ignored", reason="Unsupported extension")
            )
    if not prior:
        raise ValueError("No prior client PMR found; policy and format cannot be established")
    if not logs:
        ledger.issue(
            "ic_log_missing",
            "No investment committee log found; approval totals cannot be finalized",
            "blocker",
        )
    prior.sort(key=lambda item: item[0], reverse=True)
    if sum(p[0] == prior[0][0] for p in prior) != 1:
        raise ValueError("Ambiguous most recent prior client PMR")
    prior_text = " ".join(prior[0][2])
    abbreviation = re.search(re.escape(name) + r'\s*\(["“]([^"”]+)["”]\)', re.sub(r"\s+", " ", prior_text))
    if abbreviation and norm(abbreviation[1]) != norm(client):
        ledger.issue(
            "client_identity",
            "Prior PMR abbreviation disagrees with requested client and flash benchmark",
            "blocker",
        )
    if not abbreviation:
        ledger.issue(
            "client_alias_missing",
            "Prior PMR does not explicitly define the client abbreviation; flash benchmark code used for identity",
            "warning",
        )
    if len(deck) != 1 or len(history) != 1 or len(appendix) != 1:
        ledger.issue(
            "source_discovery", "Market deck, history or flash appendix is missing or ambiguous", "blocker"
        )
    return dict(
        root=root,
        flash=flash,
        book=book,
        name=name,
        name_id=name_id,
        prior=prior[0],
        managers=managers,
        appendix=appendix[0] if len(appendix) == 1 else None,
        logs=logs,
        history=history[0] if len(history) == 1 else None,
        deck=deck[0] if len(deck) == 1 else None,
    )
