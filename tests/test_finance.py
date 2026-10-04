"""Golden expectations live only in tests; mutations exercise unseen-run behavior."""

from pathlib import Path
import shutil

import openpyxl
import pytest
from decimal import Decimal

from pmr.evidence import Ledger
from pmr.ingest import discover, quarter, entity
from pmr.finance import build_financials, commitments
from pmr.pipeline import generate

INPUTS = Path(__file__).resolve().parents[1] / "inputs"


def calculate(root=INPUTS, client="CPERS"):
    ledger = Ledger()
    sources = discover(root, client, (2025, 4), ledger)
    data = build_financials(sources, (2025, 4), client, ledger)
    activity = commitments(sources, (2025, 4), client, data["funds"], ledger)
    return data, activity, ledger


def copied(tmp_path):
    root = tmp_path / "inputs"
    shutil.copytree(INPUTS, root)
    return root


def test_golden_financial_results():
    data, activity, ledger = calculate()
    assert not [x for x in ledger.issues if x["severity"] == "blocker"]
    assert all(x["passed"] for x in ledger.checks if x["severity"] == "blocker")
    assert any(c["delta"] == Decimal("0.51") and c["severity"] == "warning" for c in ledger.checks)
    assert data["portfolio"]["nav"] == Decimal("436116069.16")
    assert data["portfolio"]["one_net"] == pytest.approx(5.4568960063)
    assert data["portfolio"]["ex_us"] == pytest.approx(9.6199045052)
    roles = {f["name"]: f["roles"] for f in data["funds"]}
    assert roles["Redwood Logistics Trust"] == ["largest contributor", "second-largest individual position"]
    assert roles["Cornerstone Core Property Fund"] == [
        "second-largest detractor",
        "largest individual position",
    ]
    assert roles["Ironwood Value-Add Fund III"] == ["second-largest contributor"]
    assert roles["Meridian Opportunity Fund V"] == ["largest detractor"]
    assert sum(e["amount"] for e in activity["open"]) == 65_000_000
    assert len(activity["new"]) == 2
    assert {e["action"] for e in activity["reversed"]} == {"withdrawn", "lapsed"}


def test_rows_move_and_filenames_change(tmp_path):
    root = copied(tmp_path)
    path = root / "flash_4Q25.xlsx"
    book = openpyxl.load_workbook(path)
    for sheet in book:
        sheet.insert_rows(1, 4)
    book.save(path)
    path.rename(root / "unhelpful_name.xlsx")
    data, activity, ledger = calculate(root)
    assert data["portfolio"]["nav"] == Decimal("436116069.16")
    assert len(data["funds"]) == 12
    assert all(c["passed"] for c in ledger.checks if c["severity"] == "blocker")
    assert len(activity["new"]) == 2


def test_ic_order_does_not_define_chronology(tmp_path):
    root = copied(tmp_path)
    path = root / "ic_log_2025.xlsx"
    book = openpyxl.load_workbook(path)
    sheet = book.active
    rows = list(sheet.values)
    for index, row in enumerate(reversed(rows[1:]), 2):
        for column, value in enumerate(row, 1):
            sheet.cell(index, column, value)
    book.save(path)
    _, activity, _ = calculate(root)
    assert [e["name"] for e in activity["open"]] == [
        "Northgate Logistics Partners",
        "Ridgeline Value Fund IV",
    ]


def test_bad_cash_value_is_a_blocker(tmp_path):
    root = copied(tmp_path)
    path = root / "flash_4Q25.xlsx"
    book = openpyxl.load_workbook(path)
    sheet = book["CashActivity(Agg)"]
    row = next(r for r in sheet if r[0].value == "Cornerstone Core Property Fund")
    row[5].value += 1000
    book.save(path)
    _, _, ledger = calculate(root)
    assert any(x["code"] == "reconciliation" and x["severity"] == "blocker" for x in ledger.issues)


def test_fund_missing_from_cash_sheet_blocks_without_crashing(tmp_path):
    root = copied(tmp_path)
    path = root / "flash_4Q25.xlsx"
    book = openpyxl.load_workbook(path)
    sheet = book["CashActivity(Agg)"]
    row = next(r for r in sheet if r[0].value == "Blackford Special Situations Fund II")
    sheet.delete_rows(row[0].row)
    book.save(path)
    data, _, ledger = calculate(root)
    fund = next(f for f in data["funds"] if f["name"] == "Blackford Special Situations Fund II")
    assert len(data["funds"]) == 12 and fund["nav"] > 0
    assert fund["contribution"] is None and fund["ltv"] is None
    assert not data["attribution_complete"]
    assert all(
        not any("contributor" in role or "detractor" in role for role in f["roles"]) for f in data["funds"]
    )
    assert any(x["code"] == "fund_row_missing" and x["severity"] == "blocker" for x in ledger.issues)
    payload, summary = generate(root, tmp_path / "out", "CPERS", "4Q25")
    assert len(payload["data"]["activity"]["open"]) == 2
    assert payload["data"]["portfolio"]["approved_total"] == Decimal("456814526.42")
    assert payload["data"]["portfolio"]["funded_count"] == 12
    assert summary["checks_passed"] < summary["checks_total"]


def test_missing_largest_position_cash_keeps_nav_rank(tmp_path):
    root = copied(tmp_path)
    path = root / "flash_4Q25.xlsx"
    book = openpyxl.load_workbook(path)
    sheet = book["CashActivity(Agg)"]
    row = next(r for r in sheet if r[0].value == "Cornerstone Core Property Fund")
    sheet.delete_rows(row[0].row)
    book.save(path)
    payload, _ = generate(root, tmp_path / "out", "CPERS", "4Q25")
    fund = next(f for f in payload["data"]["funds"] if f["name"] == "Cornerstone Core Property Fund")
    assert fund["roles"] == ["largest individual position"]
    assert fund["contribution"] is None
    text = " ".join(b.get("text", "") for s in payload["sections"] for b in s["blocks"])
    assert "no partial performance ranking" in text


@pytest.mark.parametrize("name", ["Blackford Special Situations Fund II", "Cornerstone Core Property Fund"])
def test_blank_funded_nav_preserves_inventory_and_blocks_position_ranking(tmp_path, name):
    root = copied(tmp_path)
    path = root / "flash_4Q25.xlsx"
    book = openpyxl.load_workbook(path)
    sheet = book["FundingStatus(Agg)"]
    header = next(r for r in sheet if r[0].value == "Investment")
    column = next(c.column for c in header if c.value == "Market Value ($)")
    row = next(r for r in sheet if r[0].value == name)
    coordinate = sheet.cell(row[0].row, column).coordinate
    sheet[coordinate].value = None
    book.save(path)
    payload, summary = generate(root, tmp_path / "out", "CPERS", "4Q25")
    data = payload["data"]
    fund = next(f for f in data["funds"] if f["name"] == name)
    assert fund["nav"] is None and fund["contribution"] is not None
    assert data["portfolio"]["funded_count"] == 12
    assert len(data["activity"]["open"]) == 2
    assert all(e["name"] != name for e in data["activity"]["open"])
    assert data["portfolio"]["approved_total"] == Decimal("456814526.42")
    assert all(not any("position" in role for role in f["roles"]) for f in data["funds"])
    fact = payload["ledger"]["facts"][fund["nav_id"]]
    assert fact["value"] is None and fact["source"]["cell"] == coordinate
    assert any(i["code"] == "positions_incomplete" and i["severity"] == "blocker" for i in summary["issues"])
    assert any(i["code"] == "fund_value_missing" for i in summary["issues"])
    assert (tmp_path / "out/report.pdf").is_file()
    assert not (tmp_path / "out/approval.json").exists()


def test_fund_missing_from_returns_sheet_keeps_flash_figures(tmp_path):
    root = copied(tmp_path)
    path = root / "flash_4Q25.xlsx"
    book = openpyxl.load_workbook(path)
    sheet = book["ReturnsMultiples(Agg)"]
    row = next(r for r in sheet if r[0].value == "Blackford Special Situations Fund II")
    sheet.delete_rows(row[0].row)
    book.save(path)
    data, _, ledger = calculate(root)
    fund = next(f for f in data["funds"] if f["name"] == "Blackford Special Situations Fund II")
    assert fund["q_net"] is None and fund["contribution"] is not None
    assert any(x["code"] == "fund_row_missing" and x["severity"] == "blocker" for x in ledger.issues)


def test_conflicting_ic_amounts_block_instead_of_guessing(tmp_path):
    root = copied(tmp_path)
    path = root / "ic_log_2025.xlsx"
    book = openpyxl.load_workbook(path)
    sheet = book.active
    row = next(r for r in sheet if r[5].value == "Northgate Logistics Partners")
    row[6].value = "Approve USD 40 million to Northgate Logistics Partners, reduced from USD 50 million"
    book.save(path)
    _, activity, ledger = calculate(root)
    northgate = next(e for e in activity["new"] if e["name"] == "Northgate Logistics Partners")
    assert northgate["amount"] is None
    assert any(x["code"] == "ic_amount" and x["severity"] == "blocker" for x in ledger.issues)
    # The full pipeline must still produce a reviewable draft, not crash downstream.
    payload, summary = generate(root, tmp_path / "out", "CPERS", "4Q25")
    assert any(x["code"] == "ic_amount" for x in summary["issues"])
    assert payload["data"]["portfolio"]["approved_total"] is None
    overview = " ".join(b.get("text", "") for b in payload["sections"][0]["blocks"])
    assert "unavailable pending resolution" in overview
    assert "$0.42 billion" not in overview
    assert payload["data"]["portfolio"]["commitment"] == Decimal("391814526.42")


def test_entity_normalization_is_conservative():
    assert entity("Ironwood Value-Add Fund 3, L.P.") == entity("Ironwood Value-Add Fund III")
    assert entity("Ridgeline Value Fund IV") != entity("Ridgeline Value Fund V")
    assert entity("Kestrel U.S. Real Estate Fund") == entity("Kestrel US Real Estate Fund")


def test_document_date_wins_over_inception_quarter():
    assert quarter("Third Quarter 2025, inception 4Q 10") == (2025, 3)
    assert quarter("1Q26") == (2026, 1)


def test_unknown_committee_action_blocks_finalization(tmp_path):
    root = copied(tmp_path)
    path = root / "ic_log_2025.xlsx"
    book = openpyxl.load_workbook(path)
    sheet = book.active
    row = next(r for r in sheet if r[5].value == "Northgate Logistics Partners")
    row[9].value = "maybe approved"
    book.save(path)
    _, activity, ledger = calculate(root)
    assert len(activity["new"]) == 1
    assert any(x["code"] == "unknown_ic_action" for x in ledger.issues)


def test_zero_balance_fund_can_be_added_without_code_edits(tmp_path):
    root = copied(tmp_path)
    path = root / "flash_4Q25.xlsx"
    book = openpyxl.load_workbook(path)
    for name in ["FundingStatus(Agg)", "CashActivity(Agg)", "ReturnsMultiples(Agg)"]:
        sheet = book[name]
        row = next(r for r in sheet if r[0].value == "Tidewater Distressed Realty Fund")
        values = [c.value for c in row]
        values[0] = "Additional Real Estate Fund VII"
        for i in range(1, len(values)):
            if isinstance(values[i], (int, float)):
                values[i] = 0
        sheet.insert_rows(row[0].row + 1)
        for column, value in enumerate(values, 1):
            sheet.cell(row[0].row + 1, column, value)
    book.save(path)
    data, _, ledger = calculate(root)
    assert len(data["funds"]) == 13
    assert all(c["passed"] for c in ledger.checks if c["severity"] == "blocker")
    assert not next(f for f in data["funds"] if f["name"].startswith("Additional"))["roles"]


def test_different_client_identity_benchmark_and_policy(tmp_path):
    import fitz

    root = copied(tmp_path)
    path = root / "flash_4Q25.xlsx"
    book = openpyxl.load_workbook(path)
    old_name = "Cascadia Public Employees' Retirement System"
    new_name = "Example Workers Retirement System"
    for sheet in book:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = cell.value.replace(old_name, new_name).replace(
                        "CPERS Custom Benchmark", "EWRS Custom Benchmark"
                    )
    book.save(path)
    document = fitz.open()
    page = document.new_page()
    page.insert_text(
        (50, 50),
        new_name + "\nPerformance Measurement Report\nThird Quarter 2025\n"
        "allocated 12% of total plan assets with an allowable range of 6-22%\n"
        "<=42% per sector\n<=30% Ex-US\n<=45% / 70%",
        fontsize=10,
    )
    document.save(root / "unfamiliar_prior.pdf")
    data, _, ledger = calculate(root, "EWRS")
    assert data["portfolio"]["name"] == new_name
    assert data["portfolio"]["benchmark"] == "EWRS Custom Benchmark"
    assert data["policy"]["sector"] == [42]
    assert data["policy"]["leverage"] == [45, 70]
    assert all(c["passed"] for c in ledger.checks if c["severity"] == "blocker")


def test_after_quarter_approval_excluded(tmp_path):
    from datetime import datetime

    root = copied(tmp_path)
    path = root / "ic_log_2025.xlsx"
    book = openpyxl.load_workbook(path)
    sheet = book.active
    sheet.append(
        [
            9999,
            datetime(2026, 1, 5),
            "IC",
            "Commitment",
            "CPERS",
            "Future Fund",
            "Approve USD99 million",
            None,
            None,
            "approved",
        ]
    )
    book.save(path)
    _, activity, _ = calculate(root)
    assert all(e["name"] != "Future Fund" for e in activity["events"])
