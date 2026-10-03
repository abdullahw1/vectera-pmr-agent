"""Seeded adversarial permutations are reproducible, not unrestricted fuzzing."""
import random
import shutil
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from pmr.evidence import Ledger
from pmr.finance import build_financials, commitments
from pmr.ingest import discover

INPUTS = Path(__file__).resolve().parents[1] / "inputs"


def calculate(root):
    ledger = Ledger()
    sources = discover(root, "CPERS", (2025, 4), ledger)
    data = build_financials(sources, (2025, 4), "CPERS", ledger)
    activity = commitments(sources, (2025, 4), "CPERS", data["funds"], ledger)
    return data, activity, ledger


@pytest.mark.parametrize("seed", [7, 19, 42, 101, 503])
def test_randomized_file_row_and_ic_order_preserve_results(tmp_path, seed):
    rng = random.Random(seed)
    root = tmp_path / "inputs"
    shutil.copytree(INPUTS, root)
    path = root / "flash_4Q25.xlsx"
    book = openpyxl.load_workbook(path)
    for sheet in book:
        sheet.insert_rows(1, rng.randint(1, 12))
    book.save(path)
    path.rename(root / f"portfolio_{rng.getrandbits(32)}.xlsx")
    path = root / "ic_log_2025.xlsx"
    book = openpyxl.load_workbook(path)
    sheet = book.active
    rows = list(sheet.values)[1:]
    rng.shuffle(rows)
    for index, row in enumerate(rows, 2):
        for column, value in enumerate(row, 1):
            sheet.cell(index, column).value = value
    book.save(path)
    data, activity, ledger = calculate(root)
    assert data["portfolio"]["nav"] == Decimal("436116069.16")
    assert [e["name"] for e in activity["open"]] == ["Northgate Logistics Partners", "Ridgeline Value Fund IV"]
    assert all(c["passed"] for c in ledger.checks if c["severity"] == "blocker")


@pytest.mark.parametrize("seed", [11, 37, 83, 211, 997])
def test_randomized_currency_corruption_always_blocks(tmp_path, seed):
    rng = random.Random(seed)
    root = tmp_path / "inputs"
    shutil.copytree(INPUTS, root)
    path = root / "flash_4Q25.xlsx"
    book = openpyxl.load_workbook(path)
    sheet = book["CashActivity(Agg)"]
    names = {f["name"] for f in calculate(root)[0]["funds"]}
    row = rng.choice([r for r in sheet if r[0].value in names])
    # Source corruption is not balanced in subtotals: it must never be accepted.
    row[5].value = float(Decimal(str(row[5].value or 0)) + Decimal(rng.randint(1, 9999)) / 100)
    book.save(path)
    _, _, ledger = calculate(root)
    assert any(i["code"] == "reconciliation" and i["severity"] == "blocker" for i in ledger.issues)
