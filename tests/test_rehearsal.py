"""Full client/period adaptation, not only isolated spreadsheet parsing."""

from pathlib import Path
import os

import openpyxl

from scripts.rehearse_unseen import build_fixture, make_pdf
from pmr.ingest import pdf_pages
from pmr.pipeline import generate


def test_coherent_next_quarter_with_cross_year_approval(tmp_path, monkeypatch):
    for key in ["OPENAI_API_KEY", "ANTHROPIC_API_KEY"]:
        monkeypatch.delenv(key, raising=False)
    original = Path(__file__).resolve().parents[1] / "inputs"
    expected = build_fixture(original, tmp_path / "inputs")
    payload, summary = generate(tmp_path / "inputs", tmp_path / "output", "EWRS", "1Q26")
    data = payload["data"]
    assert data["portfolio"]["name"] == "Example Workers Retirement System"
    assert str(data["portfolio"]["nav"]) == expected["nav"]
    assert len(data["funds"]) == 13
    assert "Tidewater Distressed Realty Fund" not in {f["name"] for f in data["funds"]}
    assert set(expected["added"]) <= {f["name"] for f in data["funds"]}
    assert [e["name"] for e in data["activity"]["open"]] == expected["open_names"]
    assert any(e["name"] == "Northgate Logistics Partners" for e in data["activity"]["reversed"])
    assert not data["activity"]["new"]
    assert data["support"]["Redwood Logistics Trust"]
    assert summary["checks_passed"] == summary["checks_total"]
    assert len(payload["sections"]) == 14
    assert all(i["code"] == "chart_unavailable" for i in summary["issues"] if i["severity"] == "blocker")


def swap_columns(ws, first, second):
    for row in ws.iter_rows():
        row[first].value, row[second].value = row[second].value, row[first].value


def test_harder_unseen_layout_and_label_changes(tmp_path, monkeypatch):
    """Stack layout changes the spec allows on top of the unseen fixture; results must not move."""
    for key in ["OPENAI_API_KEY", "ANTHROPIC_API_KEY"]:
        monkeypatch.delenv(key, raising=False)
    original = Path(__file__).resolve().parents[1] / "inputs"
    build_fixture(original, tmp_path / "plain")
    build_fixture(original, tmp_path / "changed")
    root = tmp_path / "changed"

    flash = root / "renamed_current_data.xlsx"
    book = openpyxl.load_workbook(flash)
    for ws in book:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value == "EWRS Custom Benchmark":
                    # The flash no longer carries the client code; only the prior PMR links it.
                    cell.value = "Custom Benchmark (NFI-ODCE, net)"
                elif cell.value == "Apartment":
                    cell.value = "Multifamily"
                elif cell.value == "Northeast":
                    cell.value = "New England"
    cash = book["CashActivity(Agg)"]
    header = next(r for r in cash.iter_rows() if r[0].value == "Investment")
    columns = [c.value for c in header]
    swap_columns(cash, columns.index("Gross Income"), columns.index("Appreciation"))
    book.save(flash)

    redwood = root / "manager_reports" / "redwood_logistics_trust_4Q25.pdf"
    pages = [p.replace("Investments Impacting Performance", "Key Asset Drivers") for p in pdf_pages(redwood)]
    make_pdf(redwood, pages)
    (root / "notes.txt").write_text("scratch notes, not a source", encoding="utf-8")

    plain, _ = generate(tmp_path / "plain", tmp_path / "out_plain", "EWRS", "1Q26")
    changed, summary = generate(root, tmp_path / "out_changed", "EWRS", "1Q26")

    def figures(data):
        funds = {f["name"]: (str(f["contribution"]), str(f["nav"]), f["roles"]) for f in data["funds"]}
        mix = {k: [str(v["value"]) for v in data["diversification"][k]] for k in ["property", "geography"]}
        # Evidence IDs legitimately change when cells move; what the reader sees must not.
        checks = [(c["label"], c["limit"], c["actual"], c["status"]) for c in data["compliance"]]
        return funds, mix, checks, [e["name"] for e in data["activity"]["open"]]

    assert figures(changed["data"]) == figures(plain["data"])
    assert changed["data"]["portfolio"]["benchmark"] == "Custom Benchmark (NFI-ODCE, net)"
    assert changed["data"]["diversification"]["property"][0]["label"] == "Multifamily"
    assert [p["heading"] for p in changed["data"]["support"]["Redwood Logistics Trust"]] == [
        "Key Asset Drivers"
    ]
    assert summary["checks_passed"] == summary["checks_total"]
    assert all(i["code"] == "chart_unavailable" for i in summary["issues"] if i["severity"] == "blocker")
