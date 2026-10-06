"""Errors preserve honest partial drafts or a structured failure, never stale success."""

from decimal import Decimal
import json
from pathlib import Path
import shutil
import sys
from zipfile import BadZipFile

import openpyxl
import pytest

from pmr.evidence import Ledger
from pmr.finance import build_financials
from pmr.ingest import Sheet, discover, norm
from pmr.pipeline import generate
from pmr.review import approve

INPUTS = Path(__file__).resolve().parents[1] / "inputs"


@pytest.fixture(autouse=True)
def no_keys(monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.setenv("PMR_LOAD_ENV", "0")


def copied(tmp_path):
    root = tmp_path / "inputs"
    shutil.copytree(INPUTS, root)
    return root


@pytest.mark.parametrize("metric", ["Gross Income", "Manager Fees", "Appreciation", "Beginning Market Value ($)", "Market Value ($)", "LTV"])
def test_blank_fund_cash_figure_creates_a_blocked_draft(tmp_path, metric):
    root = copied(tmp_path)
    path = root / "flash_4Q25.xlsx"
    book = openpyxl.load_workbook(path)
    sheet = Sheet(path, book["CashActivity(Agg)"], Ledger())
    _, headers = sheet.headers()
    row = sheet.find("Cornerstone Core Property Fund", numeric=True)
    cell = row[sheet.column(headers, metric)]
    cell.value = None
    coordinate = cell.coordinate
    book.save(path)
    output = tmp_path / "output"
    payload, summary = generate(root, output, "CPERS", "4Q25")
    facts = payload["ledger"]["facts"]
    assert any(f["value"] is None and f["source"].get("cell") == coordinate
               and f["source"].get("sheet") == "CashActivity(Agg)" for f in facts.values())
    assert any(i["code"] == "fund_value_missing" and i["severity"] == "blocker" for i in summary["issues"])
    fund = next(f for f in payload["data"]["funds"] if f["name"] == "Cornerstone Core Property Fund")
    if metric in {"Gross Income", "Manager Fees", "Appreciation"}:
        assert fund["contribution"] is None
        assert not payload["data"]["attribution_complete"]
        assert all(not any("contributor" in r or "detractor" in r for r in f["roles"]) for f in payload["data"]["funds"])
    assert (output / "report.pdf").exists()
    with pytest.raises(ValueError, match="Finalization blocked"):
        approve(output, dict(draft_hash=payload["draft_hash"], reviewer="Automated test"))


def test_blank_capital_flow_still_means_zero(tmp_path):
    root = copied(tmp_path)
    path = root / "flash_4Q25.xlsx"
    book = openpyxl.load_workbook(path)
    sheet = Sheet(path, book["CashActivity(Agg)"], Ledger())
    _, headers = sheet.headers()
    row = sheet.find("Cornerstone Core Property Fund", numeric=True)
    row[sheet.column(headers, "Contributions")].value = None
    book.save(path)
    ledger = Ledger()
    data = build_financials(discover(root, "CPERS", (2025, 4), ledger), (2025, 4), "CPERS", ledger)
    fund = next(f for f in data["funds"] if f["name"] == "Cornerstone Core Property Fund")
    assert fund["contributions"] == Decimal(0)
    assert ledger.facts[fund["contributions_id"]]["source"]["interpretation"] == "blank capital flow = zero"


def test_blank_fund_return_has_a_source_cell():
    ledger = Ledger()
    data = build_financials(discover(INPUTS, "CPERS", (2025, 4), ledger), (2025, 4), "CPERS", ledger)
    fund = next(f for f in data["funds"] if f["name"] == "Tidewater Distressed Realty Fund")
    assert fund["q_net"] is None
    fact = ledger.facts[fund["q_net_id"]]
    assert fact["source"]["sheet"] == "ReturnsMultiples(Agg)"
    assert fact["source"]["cell"] and fact["source"]["status"] == "unavailable"


@pytest.mark.parametrize("name", ["CashActivity(Agg)", "AnnReturns(Agg)", "All Geographic"])
def test_missing_essential_sheet_names_the_file_and_sheet(tmp_path, name):
    root = copied(tmp_path)
    path = root / "flash_4Q25.xlsx"
    book = openpyxl.load_workbook(path)
    del book[name]
    book.save(path)
    ledger = Ledger()
    with pytest.raises(ValueError, match="Required sheets missing in flash_4Q25.xlsx") as error:
        build_financials(discover(root, "CPERS", (2025, 4), ledger), (2025, 4), "CPERS", ledger)
    assert name in str(error.value)


@pytest.mark.parametrize("label", ["Allocation", "Total Plan Assets", "Performance Summary"])
def test_missing_composition_label_has_a_useful_error(tmp_path, label):
    root = copied(tmp_path)
    path = root / "flash_4Q25.xlsx"
    book = openpyxl.load_workbook(path)
    for row in book["FundingStatus(Agg)"]:
        for cell in row:
            if norm(cell.value) == norm(label):
                cell.value = "Unrecognized source label"
    book.save(path)
    ledger = Ledger()
    with pytest.raises(ValueError) as error:
        build_financials(discover(root, "CPERS", (2025, 4), ledger), (2025, 4), "CPERS", ledger)
    assert "flash_4Q25.xlsx/FundingStatus(Agg)" in str(error.value)
    assert label in str(error.value)


@pytest.mark.parametrize("value", [True, float("nan"), float("inf"), 1e100, 1.001, "95000"])
def test_invalid_money_always_reports_its_location(value):
    book = openpyxl.Workbook()
    sheet = Sheet(Path("flash.xlsx"), book.active, Ledger())
    book.active["B2"] = value
    with pytest.raises(ValueError, match=r"flash.xlsx/Sheet/B2"):
        sheet.number(book.active["B2"], monetary=True, metric="Market Value")


def test_duplicate_columns_are_not_silently_overwritten():
    book = openpyxl.Workbook()
    book.active.append(["Investment", "Market Value ($)", "Market Value ($)"])
    with pytest.raises(ValueError, match="Ambiguous column"):
        Sheet(Path("flash.xlsx"), book.active, Ledger()).headers()


def test_missing_or_ambiguous_prefix_column_is_explicit():
    sheet = Sheet(Path("history.xlsx"), openpyxl.Workbook().active, Ledger())
    with pytest.raises(ValueError, match="matched 0 columns"):
        sheet.column({}, "target allocation", prefix=True)
    with pytest.raises(ValueError, match="matched 2 columns"):
        sheet.column({"target allocation old": 1, "target allocation new": 2}, "target allocation", prefix=True)


@pytest.mark.parametrize("error,exit_code,category", [
    (BadZipFile("File is not a zip file"), 2, "source_data"),
    (PermissionError("Access denied"), 2, "filesystem"),
    (RuntimeError("secret-must-not-appear-in-diagnostics"), 3, "unexpected"),
    (StopIteration(), 2, "source_data"),
])
def test_cli_records_failures_without_unhandled_tracebacks(tmp_path, monkeypatch, capsys, error, exit_code, category):
    import pmr.__main__ as cli

    def fail(*args):
        raise error

    monkeypatch.setattr(cli, "generate", fail)
    monkeypatch.setattr(sys, "argv", ["pmr", "generate", "--client", "CPERS", "--quarter", "4Q25", "--output", str(tmp_path)])
    with pytest.raises(SystemExit) as stopped:
        cli.main()
    assert stopped.value.code == exit_code
    failure = json.loads((tmp_path / "failure.json").read_text())
    assert failure["status"] == "blocked" and failure["category"] == category
    assert failure["error_action"] and failure["error_type"] == type(error).__name__
    assert "secret-must-not-appear" not in json.dumps(failure)
    assert "secret-must-not-appear" not in capsys.readouterr().err


def test_failed_rerun_removes_previous_draft_and_approval(tmp_path):
    output = tmp_path / "output"
    generate(INPUTS, output, "CPERS", "4Q25")
    (output / "approval.json").write_text("old approval")
    (output / "final_report.pdf").write_bytes(b"old final report")
    with pytest.raises(ValueError, match="found 0"):
        generate(INPUTS, output, "WRONGCLIENT", "4Q25")
    for name in ("report.pdf", "draft.json", "review.html", "manifest.json", "verification.json", "approval.json", "final_report.pdf", "CPERS_PMR_4Q25_DRAFT.pdf"):
        assert not (output / name).exists(), name


def test_output_cannot_modify_the_source_folder(tmp_path):
    root = copied(tmp_path)
    original = (root / "flash_4Q25.xlsx").read_bytes()
    with pytest.raises(ValueError, match="Output folder must be separate"):
        generate(root, root / "generated", "CPERS", "4Q25")
    assert (root / "flash_4Q25.xlsx").read_bytes() == original


@pytest.mark.parametrize("filename", ["flash_4Q25.xlsx", "market_outlook_4Q25.pptx"])
def test_corrupt_documents_produce_a_structured_cli_failure(tmp_path, monkeypatch, filename):
    import pmr.__main__ as cli

    root = copied(tmp_path)
    (root / filename).write_bytes(b"damaged document")
    output = tmp_path / "output"
    monkeypatch.setattr(sys, "argv", ["pmr", "generate", "--client", "CPERS", "--quarter", "4Q25",
                                      "--inputs", str(root), "--output", str(output)])
    if filename.endswith(".pptx"):
        cli.main()
        payload = json.loads((output / "draft.json").read_text())
        assert any(i["code"] == "unreadable_input" and i["severity"] == "blocker" for i in payload["ledger"]["issues"])
        assert (output / "report.pdf").exists()
        with pytest.raises(ValueError, match="Finalization blocked"):
            approve(output, dict(draft_hash=payload["draft_hash"], reviewer="Automated test"))
        return
    with pytest.raises(SystemExit) as stopped:
        cli.main()
    assert stopped.value.code == 2
    failure = json.loads((output / "failure.json").read_text())
    assert failure["error_title"] == "Source document is damaged or unreadable"
    assert not (output / "report.pdf").exists()
    assert any(s["status"] == "failed" for s in json.loads((output / "diagnostics.json").read_text())["stages"])


@pytest.mark.parametrize("sheet_name,label,metric,grouped", [
    ("FundingStatus(Agg)", "Vectera Initiated Investments", "Market Value ($)", False),
    ("FundingStatus(Agg)", "Vectera Initiated Investments", "Commitment Amount", False),
    ("CashActivity(Agg)", "Vectera Initiated Investments", "Gross Income", False),
    ("CashActivity(Agg)", "Strategic Investments", "LTV", False),
    ("ReturnsMultiples(Agg)", "CPERS Custom Benchmark", "Quarter NET", True),
    ("AnnReturns(Agg)", "CPERS Custom Benchmark", "1 Year Net", True),
])
def test_missing_load_bearing_metric_stops_with_a_source_locator(tmp_path, sheet_name, label, metric, grouped):
    root = copied(tmp_path)
    path = root / "flash_4Q25.xlsx"
    book = openpyxl.load_workbook(path)
    sheet = Sheet(path, book[sheet_name], Ledger())
    _, headers = sheet.headers(grouped=grouped)
    row = sheet.find(label, numeric=True)
    cell = row[sheet.column(headers, metric)]
    cell.value = None
    coordinate = cell.coordinate
    book.save(path)
    ledger = Ledger()
    with pytest.raises(ValueError) as error:
        build_financials(discover(root, "CPERS", (2025, 4), ledger), (2025, 4), "CPERS", ledger)
    assert f"flash_4Q25.xlsx/{sheet_name}/{coordinate}" in str(error.value)


def test_zero_target_stops_without_division_by_zero(tmp_path):
    root = copied(tmp_path)
    path = root / "flash_4Q25.xlsx"
    book = openpyxl.load_workbook(path)
    sheet = book["FundingStatus(Agg)"]
    label = next(c for row in sheet for c in row if norm(c.value) == "allocation")
    sheet.cell(label.row + 1, label.column).value = 0
    book.save(path)
    with pytest.raises(ValueError, match="Required allocation ratios.*target at flash_4Q25.xlsx"):
        generate(root, tmp_path / "output", "CPERS", "4Q25")
    assert not (tmp_path / "output/report.pdf").exists()


def test_late_failure_removes_the_just_rendered_report(tmp_path, monkeypatch):
    def fail(*args):
        raise RuntimeError("Unexpected verification failure")

    monkeypatch.setattr("pmr.pipeline.write_manifests", fail)
    output = tmp_path / "output"
    with pytest.raises(RuntimeError):
        generate(INPUTS, output, "CPERS", "4Q25")
    assert not (output / "report.pdf").exists()
    assert not (output / "draft.json").exists()
    assert (output / "progress.json").exists()


def test_cli_never_writes_failure_records_into_inputs(tmp_path, monkeypatch):
    import pmr.__main__ as cli

    root = copied(tmp_path)
    before = sorted(str(p.relative_to(root)) for p in root.rglob("*") if p.is_file())
    monkeypatch.setattr(sys, "argv", ["pmr", "generate", "--client", "CPERS", "--quarter", "4Q25",
                                      "--inputs", str(root), "--output", str(root)])
    with pytest.raises(SystemExit):
        cli.main()
    assert sorted(str(p.relative_to(root)) for p in root.rglob("*") if p.is_file()) == before


def test_cli_redacts_configured_keys_from_validation_errors(tmp_path, monkeypatch, capsys):
    import pmr.__main__ as cli

    secret = "example-private-key-do-not-persist"
    monkeypatch.setenv("OPENAI_API_KEY", secret)

    def fail(*args):
        raise ValueError("Bad response for " + secret)

    monkeypatch.setattr(cli, "generate", fail)
    monkeypatch.setattr(sys, "argv", ["pmr", "generate", "--client", "CPERS", "--quarter", "4Q25", "--output", str(tmp_path)])
    with pytest.raises(SystemExit):
        cli.main()
    assert secret not in (tmp_path / "failure.json").read_text()
    assert secret not in capsys.readouterr().err


def test_missing_annual_layout_blocks_report_without_a_rendering_crash(tmp_path):
    from pmr.report import build_sections

    ledger = Ledger()
    sources = discover(INPUTS, "CPERS", (2025, 4), ledger)
    data = build_financials(sources, (2025, 4), "CPERS", ledger)
    data["portfolio"]["annual"] = []
    data["activity"] = {}
    with pytest.raises(ValueError, match="Required annualized chart horizons.*prior_pmr_3Q25.pdf"):
        build_sections(data, sources, ledger)
